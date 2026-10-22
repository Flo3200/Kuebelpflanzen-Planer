"""Round 2, packages O and P: every plant can be seen as a picture (photo first, else an AI image),
and the spreadsheet-like tables with export."""

from __future__ import annotations

import io

from openpyxl import load_workbook
from PIL import Image

from tests.test_phase3_estimates import make_order
from tests.test_round2_inbox import MESSAGE, _photos


def test_ai_image_for_plants_without_photo(new_tenant):
    _, owner = new_tenant()
    o = {"id": make_order(owner, plants=2)}
    detail = owner.get(f"/api/orders/{o['id']}").json()
    assert all(p["image"] is None for p in detail["plants"])
    r = owner.post(f"/api/orders/{o['id']}/plant-images", json={"sync": True}).json()
    assert r["status"] == "done" and r["result"]["images"] == 2
    plants = owner.get(f"/api/orders/{o['id']}").json()["plants"]
    img = plants[0]["image"]
    assert img["generated"] is True
    raw = owner.get(img["url"]).content
    with Image.open(io.BytesIO(raw)) as im:
        assert im.size == (600, 800)
    # a second run does not draw again; "force" replaces the AI image (still one per plant)
    assert (
        owner.post(f"/api/orders/{o['id']}/plant-images", json={"sync": True}).json()["result"]["images"] == 0
    )
    owner.post(f"/api/plants/{plants[0]['id']}/image", json={"force": True})
    photos = owner.get(f"/api/orders/{o['id']}").json()["photos"]
    assert sum(1 for p in photos if p["kind"] == "generated" and p["plant_id"] == plants[0]["id"]) == 1


def test_real_photo_wins_and_confirm_makes_images(new_tenant):
    _, owner = new_tenant()
    d = owner.post(
        "/api/drafts", expect=201, data={"text": MESSAGE, "sync": "true"}, files=_photos(2, 0)
    ).json()
    res = owner.post(f"/api/drafts/{d['id']}/confirm", json={}).json()
    from app.jobs.queue import run_pending

    run_pending()  # the image job queued by the confirmation
    plants = owner.get(f"/api/orders/{res['order_id']}").json()["plants"]
    by_name = {p["species_name"]: p["image"] for p in plants}
    assert by_name["Oleander"]["generated"] is False  # has the pasted photos
    assert by_name["Olivenbaum"]["generated"] is False


def test_tables_and_export(new_tenant):
    _, owner = new_tenant()
    o = {"id": make_order(owner, plants=3)}
    owner.post(f"/api/orders/{o['id']}/plant-images", json={"sync": True})
    plants = owner.get("/api/tables/plants").json()
    keys = {c["key"] for c in plants["columns"]}
    assert {"image", "species_name", "street", "city", "customer"} <= keys
    assert len(plants["rows"]) == 3
    row = plants["rows"][0]
    assert row["street"] == "Teststraße 1" and row["image"]["generated"] is True
    orders = owner.get("/api/tables/orders").json()["rows"]
    assert orders[0]["n_plants"] == 3 and len(orders[0]["plant_images"]) == 3
    customers = owner.get("/api/tables/customers").json()["rows"]
    assert customers[0]["n_plants"] == 3
    owner.get("/api/tables/stops")
    owner.get("/api/tables/unknown", expect=400)
    # editing a cell goes through the normal endpoint
    owner.patch(f"/api/plants/{row['id']}", json={"total_height_cm": 222})
    assert owner.get("/api/tables/plants").json()["rows"][0]["total_height_cm"] == 222
    # export with the chosen columns
    x = owner.get("/api/tables/plants/export", params={"cols": "number,species_name,city,image"})
    wb = load_workbook(io.BytesIO(x.content))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["Nr.", "Pflanze", "Ort"]  # pictures are not exported
    assert ws.max_row == 4
    csv = owner.get("/api/tables/orders/export", params={"fmt": "csv", "cols": "number,city"}).content.decode(
        "utf-8-sig"
    )
    assert csv.splitlines()[0] == "Nr.;Ort"
