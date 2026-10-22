"""Tables (round 2, package P): spreadsheet-like overviews. One flat row per plant, order, customer or
tour stop with every column the page can show; the page lets the user tick the columns, sort, filter,
group, edit cells (through the normal endpoints) and export to Excel/CSV."""

from __future__ import annotations

import csv
import io
import uuid
from typing import Any

from fastapi import APIRouter
from fastapi.responses import Response
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.ai.images import plant_image
from app.api.deps import DB, CurrentActor
from app.core.context import Actor
from app.core.errors import AppError
from app.i18n import t
from app.models.masterdata import StaffMember, Vehicle
from app.models.orders import Address, Customer, PickupOrder, Plant
from app.models.planning import Plan, Tour, TourStop

router = APIRouter(prefix="/api/tables", tags=["tables"])

# column: (key, type, editable field or None, default visible)
# types: text | number | bool | date | image | images | status
COLUMNS: dict[str, list[tuple[str, str, str | None, bool]]] = {
    "plants": [
        ("image", "image", None, True),
        ("number", "number", None, True),
        ("species_name", "text", "species_name", True),
        ("label", "text", "label", False),
        ("customer", "text", None, True),
        ("street", "text", None, True),
        ("postal_code", "text", None, False),
        ("city", "text", None, True),
        ("order_number", "text", None, False),
        ("order_status", "status", None, False),
        ("total_height_cm", "number", "total_height_cm", True),
        ("pot_diameter_cm", "number", "pot_diameter_cm", True),
        ("crown_width_cm", "number", "crown_width_cm", False),
        ("weight_kg", "number", "weight_kg", True),
        ("weight_measured_kg", "number", None, False),
        ("can_tilt", "bool", "can_tilt", False),
        ("frost_sensitive", "bool", "frost_sensitive", False),
        ("sensitive", "bool", "sensitive", False),
        ("team_size", "number", None, False),
        ("tour_date", "date", None, True),
        ("vehicle", "text", None, False),
        ("notes", "text", "notes", False),
    ],
    "orders": [
        ("plant_images", "images", None, True),
        ("number", "text", None, True),
        ("customer", "text", None, True),
        ("street", "text", None, True),
        ("postal_code", "text", None, False),
        ("city", "text", None, True),
        ("status", "status", None, True),
        ("direction", "text", None, False),
        ("n_plants", "number", None, True),
        ("plant_names", "text", None, True),
        ("weight_kg", "number", None, False),
        ("earliest_date", "date", "earliest_date", False),
        ("latest_date", "date", "latest_date", False),
        ("duration_min", "number", None, False),
        ("team_size", "number", None, False),
        ("tour_date", "date", None, True),
        ("vehicle", "text", None, False),
        ("notes", "text", "notes", False),
    ],
    "customers": [
        ("name", "text", "name", True),
        ("company", "text", "company", False),
        ("phone", "text", "phone", False),
        ("email", "text", "email", False),
        ("street", "text", None, True),
        ("postal_code", "text", None, True),
        ("city", "text", None, True),
        ("access_notes", "text", None, False),
        ("n_orders", "number", None, True),
        ("n_plants", "number", None, True),
        ("plant_images", "images", None, False),
        ("last_activity_at", "date", None, False),
    ],
    "stops": [
        ("tour_date", "date", None, True),
        ("vehicle", "text", None, True),
        ("team", "text", None, True),
        ("position", "number", None, True),
        ("arrival", "text", None, True),
        ("customer", "text", None, True),
        ("street", "text", None, True),
        ("city", "text", None, True),
        ("plant_images", "images", None, True),
        ("plant_names", "text", None, True),
        ("n_plants", "number", None, False),
        ("team_size", "number", None, False),
        ("predicted_min", "number", None, False),
        ("status", "status", None, True),
    ],
}
SOURCES = tuple(COLUMNS)


def _hm(m: int | None) -> str | None:
    return None if m is None else f"{m // 60:02d}:{m % 60:02d}"


def _active_stops(db: Session, tid: uuid.UUID) -> dict[uuid.UUID, tuple[TourStop, Tour]]:
    """Latest stop of every order in the active plans."""
    rows = db.execute(
        select(TourStop, Tour)
        .join(Tour, Tour.id == TourStop.tour_id)
        .join(Plan, Plan.id == Tour.plan_id)
        .where(Plan.tenant_id == tid, Plan.is_active.is_(True), Plan.kind == "plan")
        .order_by(Tour.date)
    ).all()
    return {s.order_id: (s, tour) for s, tour in rows}


def _names(db: Session, tid: uuid.UUID) -> tuple[dict[uuid.UUID, str], dict[str, str]]:
    vehicles = {v.id: v.name for v in db.scalars(select(Vehicle).where(Vehicle.tenant_id == tid))}
    staff = {str(s.id): s.name for s in db.scalars(select(StaffMember).where(StaffMember.tenant_id == tid))}
    return vehicles, staff


def build_rows(db: Session, actor: Actor, source: str) -> list[dict[str, Any]]:
    tid = actor.require_tenant()
    if source not in SOURCES:
        raise AppError("errors.validation")
    orders = {
        o.id: o
        for o in db.scalars(
            select(PickupOrder).where(PickupOrder.tenant_id == tid, PickupOrder.deleted_at.is_(None))
        )
    }
    customers = {
        c.id: c
        for c in db.scalars(select(Customer).where(Customer.tenant_id == tid, Customer.deleted_at.is_(None)))
    }
    addresses = {
        a.id: a
        for a in db.scalars(select(Address).where(Address.tenant_id == tid, Address.deleted_at.is_(None)))
    }
    plants_by_order: dict[uuid.UUID, list[Plant]] = {}
    for p in db.scalars(
        select(Plant).where(Plant.tenant_id == tid, Plant.deleted_at.is_(None)).order_by(Plant.number)
    ):
        plants_by_order.setdefault(p.order_id, []).append(p)
    stops = _active_stops(db, tid)
    vehicles, staff = _names(db, tid)
    image_cache: dict[uuid.UUID, dict[str, Any] | None] = {}

    def img(p: Plant) -> dict[str, Any] | None:
        if p.id not in image_cache:
            image_cache[p.id] = plant_image(db, p)
        return image_cache[p.id]

    def plants_of(o: PickupOrder) -> list[Plant]:
        return plants_by_order.get(o.return_of_id or o.id, [])

    def addr(o: PickupOrder | None) -> dict[str, Any]:
        a = addresses.get(o.address_id) if o else None
        return {
            "street": a.street if a else None,
            "postal_code": a.postal_code if a else None,
            "city": a.city if a else None,
            "access_notes": a.access_notes if a else None,
        }

    def tour_of(o: PickupOrder) -> dict[str, Any]:
        st = stops.get(o.id)
        if st is None:
            return {"tour_date": None, "vehicle": None}
        _, tour = st
        return {"tour_date": tour.date, "vehicle": vehicles.get(tour.vehicle_id)}

    def images_of(plants: list[Plant]) -> list[dict[str, Any]]:
        return [x for x in (img(p) for p in plants[:6]) if x]

    rows: list[dict[str, Any]] = []
    if source == "plants":
        for o in orders.values():
            c = customers.get(o.customer_id)
            crew = {p["id"]: p for p in (o.estimate or {}).get("plants", [])}
            for p in plants_of(o) if o.direction == "pickup" else []:
                rows.append(
                    {
                        "id": p.id,
                        "order_id": o.id,
                        "image": img(p),
                        "number": p.number,
                        "species_name": p.species_name,
                        "label": p.label,
                        "customer": c.name if c else None,
                        **addr(o),
                        "order_number": o.number,
                        "order_status": o.status,
                        "total_height_cm": p.total_height_cm,
                        "pot_diameter_cm": p.pot_diameter_cm,
                        "crown_width_cm": p.crown_width_cm,
                        "weight_kg": p.weight_kg,
                        "weight_measured_kg": p.weight_measured_kg,
                        "can_tilt": p.can_tilt,
                        "frost_sensitive": p.frost_sensitive,
                        "sensitive": p.sensitive,
                        "team_size": (crew.get(str(p.id)) or {}).get("crew"),
                        **tour_of(o),
                        "notes": p.notes,
                        "estimated": list(p.estimated_fields or []),
                    }
                )
        rows.sort(key=lambda r: r["number"] or 0)
    elif source == "orders":
        for o in sorted(orders.values(), key=lambda x: x.number):
            c = customers.get(o.customer_id)
            plants = plants_of(o)
            est = o.estimate or {}
            rows.append(
                {
                    "id": o.id,
                    "order_id": o.id,
                    "plant_images": images_of(plants),
                    "number": o.number,
                    "customer": c.name if c else None,
                    **addr(o),
                    "status": o.status,
                    "direction": t(f"enums.{o.direction}", _default=o.direction),
                    "n_plants": len(plants),
                    "plant_names": ", ".join(p.species_name or p.label or "?" for p in plants),
                    "weight_kg": round(sum((p.weight_measured_kg or p.weight_kg or 0) for p in plants), 1),
                    "earliest_date": o.earliest_date,
                    "latest_date": o.latest_date,
                    "duration_min": est.get("duration_min"),
                    "team_size": est.get("team_size"),
                    **tour_of(o),
                    "notes": o.notes,
                }
            )
    elif source == "customers":
        by_customer: dict[uuid.UUID, list[PickupOrder]] = {}
        for o in orders.values():
            by_customer.setdefault(o.customer_id, []).append(o)
        for c in sorted(customers.values(), key=lambda x: x.name.lower()):
            os_ = by_customer.get(c.id, [])
            a = next((addresses[o.address_id] for o in os_ if o.address_id in addresses), None)
            if a is None:
                a = next((x for x in addresses.values() if x.customer_id == c.id), None)
            plants = [p for o in os_ if o.direction == "pickup" for p in plants_of(o)]
            rows.append(
                {
                    "id": c.id,
                    "customer_id": c.id,
                    "name": c.name,
                    "company": c.company,
                    "phone": c.phone,
                    "email": c.email,
                    "street": a.street if a else None,
                    "postal_code": a.postal_code if a else None,
                    "city": a.city if a else None,
                    "access_notes": a.access_notes if a else None,
                    "n_orders": len(os_),
                    "n_plants": len(plants),
                    "plant_images": images_of(plants),
                    "last_activity_at": c.last_activity_at,
                }
            )
    else:  # stops
        for order_id, (s, tour) in stops.items():
            o = orders.get(order_id)
            if o is None:
                continue
            c = customers.get(o.customer_id)
            plants = plants_of(o)
            rows.append(
                {
                    "id": s.id,
                    "order_id": o.id,
                    "tour_id": tour.id,
                    "tour_date": tour.date,
                    "vehicle": vehicles.get(tour.vehicle_id),
                    "team": ", ".join(staff.get(str(m), "?") for m in tour.member_ids or []),
                    "position": s.position + 1,
                    "arrival": _hm(s.arrival_min),
                    "customer": c.name if c else None,
                    **addr(o),
                    "plant_images": images_of(plants),
                    "plant_names": ", ".join(p.species_name or p.label or "?" for p in plants),
                    "n_plants": len(plants),
                    "team_size": s.team_size,
                    "predicted_min": round(s.predicted_min or 0),
                    "status": s.status,
                }
            )
        rows.sort(key=lambda r: (str(r["tour_date"]), r["vehicle"] or "", r["position"]))
    return rows


def columns(source: str) -> list[dict[str, Any]]:
    return [
        {
            "key": key,
            "type": kind,
            "label": t(f"tables.columns.{key}", _default=key),
            "edit": field,
            "default": default,
        }
        for key, kind, field, default in COLUMNS[source]
    ]


@router.get("/{source}")
def table(source: str, actor: CurrentActor, db: DB) -> dict[str, Any]:
    actor.require("orders.read")
    if source == "customers":
        actor.require("customers.read")
    rows = build_rows(db, actor, source)
    return {"source": source, "columns": columns(source), "rows": rows}


def _cell(v: Any) -> Any:
    if isinstance(v, dict):
        return "KI-Bild" if v.get("generated") else "Foto"
    if isinstance(v, list):
        return f"{len(v)} Bilder"
    if isinstance(v, bool):
        return "ja" if v else "nein"
    return "" if v is None else v


@router.get("/{source}/export")
def export(source: str, actor: CurrentActor, db: DB, cols: str = "", fmt: str = "xlsx") -> Response:
    """Export of the chosen columns (comma-separated keys) as Excel file or CSV."""
    actor.require("orders.read")
    if source == "customers":
        actor.require("customers.read")
    rows = build_rows(db, actor, source)
    defs = {c["key"]: c for c in columns(source)}
    keys = [k for k in cols.split(",") if k in defs] or [c["key"] for c in columns(source) if c["default"]]
    keys = [k for k in keys if defs[k]["type"] not in ("image", "images")] or keys
    head = [defs[k]["label"] for k in keys]
    name = f"{t(f'tables.sources.{source}', _default=source)}"
    from urllib.parse import quote

    def disposition(ext: str) -> str:
        ascii_name = name.encode("ascii", "ignore").decode() or source
        return f"attachment; filename=\"{ascii_name}.{ext}\"; filename*=UTF-8''{quote(name)}.{ext}"

    if fmt == "csv":
        out = io.StringIO()
        w = csv.writer(out, delimiter=";")
        w.writerow(head)
        for r in rows:
            w.writerow([_cell(r.get(k)) for k in keys])
        return Response(
            ("﻿" + out.getvalue()).encode("utf-8"),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": disposition("csv")},
        )
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = name[:30]
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([_cell(r.get(k)) for k in keys])
    for i, k in enumerate(keys, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(10, min(40, len(defs[k]["label"]) + 6))
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": disposition("xlsx")},
    )
